# -*- coding: utf-8 -*-
"""用例表规则自检（不需要手机）：验证 tools/export_cases.py 的校验与体检规则。

为什么要它：这些规则是"表格填错时唯一的护栏" —— 但护栏本身没人验过。
真机上的行为有真机跑兜底，纯逻辑的规则就用这份自检兜底：每条规则都要
**正反两个方向**都验一遍（合规的表不报错；违规的表必须报出预期的那句话）。

用法：python scripts/selftest-export-cases.py
退出码：0 = 全部通过；1 = 有不符合预期的
"""

import os
import re
import sys
import tempfile

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "tools"))

from openpyxl import Workbook  # noqa: E402

import export_cases as m  # noqa: E402

TMP = tempfile.mkdtemp(prefix="uitest_selftest_")
PATH = os.path.join(TMP, "用例表.xlsx")

failed = []
passed = 0


def make(rows, header=None, sheet="用例"):
    wb = Workbook()
    ws = wb.active
    ws.title = sheet
    ws.append(header or m.HEADERS)
    for r in rows:
        ws.append(r)
    wb.save(PATH)
    return PATH


def ok(name, rows, want_warn_count=0, **kw):
    """正例：应当没有任何错误；并检查提示条数"""
    global passed
    try:
        cases = m.validate(m.read_rows(make(rows, **kw)))
        warns = m.lint_cases(cases)
    except m.TableError as e:
        failed.append("%s：本该通过，却报错 %s" % (name, e.errors))
        return
    if len(warns) != want_warn_count:
        failed.append("%s：提示条数 %d，期望 %d（%s）" % (name, len(warns), want_warn_count, warns))
        return
    passed += 1
    print("OK   正例 %s" % name)


def bad_case(name, rows, want, **kw):
    """反例（**归得到某条用例身上**）：导出照常成功，但那条用例自己带着错误。

    这是 2026/10/12 owner 定的口径：表格里有一条填错，不该废掉整轮测试 ——
    错误归到那条用例，真机跑的时候只有它红。所以这里断言的是"错误挂在用例上"，
    而不是"整张表报错"。
    """
    global passed
    try:
        cases = m.validate(m.read_rows(make(rows, **kw)))
    except m.TableError as e:
        failed.append("%s：本该把错误归到用例身上，却整张表报错 %s" % (name, e.errors))
        return
    msgs = [msg for c in cases for msg in (c.get("errors") or [])]
    if any(want in msg for msg in msgs):
        passed += 1
        print("OK   反例 %s -> %s" % (name, msgs[0][:60]))
        return
    failed.append("%s：用例上没带着「%s」这个错误，实际是 %s" % (name, want, msgs or "（一条错误都没有）"))


def bad_table(name, rows, want, **kw):
    """反例（**归不到用例身上**）：整张表就有问题，必须报错、不导出"""
    global passed
    try:
        m.validate(m.read_rows(make(rows, **kw)))
    except m.TableError as e:
        msg = " / ".join(e.errors)
        if want in msg:
            passed += 1
            print("OK   反例 %s -> %s" % (name, e.errors[0][:60]))
            return
        failed.append("%s：报错里没有「%s」，实际是 %s" % (name, want, msg))
        return
    failed.append("%s：本该报错（%s）却通过了" % (name, want))


GOOD = ["回到首页", "", "检查文字显示", "积分商城", "", "", ""]

# ---------- 正例 ----------
ok(
    "合规用例（起点 + 检查 + 交还）",
    [
        ["A_001", "合规", "手机在小程序首页", 1, "回到首页", "", "", "只读"],
        ["A_001", "", "", 2, "检查文字显示", "积分商城", "", ""],
        ["A_001", "", "", 3, "回到首页", "", "", ""],
    ],
)
ok("只有一步的最小用例", [["A_002", "最小", "", 1, "回到首页", "", "", ""]], want_warn_count=1)  # 没有检查 -> 1 条提示

# ---------- 校验反例（填错就归到那条用例身上，导出照常成功）----------
# 注意：这些"填错"都**不会**让整张表报废 —— 真机跑的时候只有那一条红。
bad_case("不认识的动作", [["B_001", "名字", "", 1, "打开首贝", "", "", ""]], "不认识动作")
bad_case("动作缺参数1", [["B_002", "名字", "", 1, "打开搜索并输入", "", "", ""]], "参数1")
bad_case("多余参数2", [["B_003", "名字", "", 1, "回到首页", "", "多填的", ""]], "用不到「参数2」")
bad_case("固定候选项写错", [["B_004", "名字", "", 1, "切换首页筛选栏", "最新问题", "", ""]], "只能填这 4 个之一")
# 2026/10/14：滑动参数2 改成"方向可带幅度"（上 / 下 / 上 30%），报错文案也跟着换了
bad_case("滑动方向写错", [["B_005", "名字", "", 1, "滑动", "1", "左", ""]], "参数2 只能填")
bad_case(
    "步骤序号重复",
    [["B_006", "名字", "", 1, "回到首页", "", "", ""], ["B_006", "", "", 1, "返回上一页", "", "", ""]],
    "重复",
)
bad_case("动作那一格是空的", [["B_007", "名字", "", 1, "", "", "", ""]], "「动作」是空的")
bad_case("没写用例标题", [["B_011", "", "", 1, "回到首页", "", "", ""]], "没写「用例标题」")

# ---------- 按 id 定位：动作已删，另一条路（选择器参数）也要挡（2026/10/14）----------
# 为什么盯它：id 会变（重进一次页面就全换），按它定位的用例今天绿明天红。
# 三个按 id 定位的动作已经删了，这里验的是"把 id 抄成选择器"那条路也堵上了。
bad_case(
    "选择器参数里写 #id",
    [["B_012", "名字", "", 1, "点选择器", "#_AQ", "", ""]],
    "按 id 定位",
)
bad_case(
    "选择器参数里写 [id=...]",
    [["B_013", "名字", "", 1, "检查元素存在", '[id="_AQ"]', "", ""]],
    "按 id 定位",
)


def id_selector_rule():
    """纯函数正反例：该认的认出来，**合法选择器一个都不许误伤**（误伤比漏判更坏 ——
    填表的人会以为"连 .xxx 都不能用"）。"""
    global passed
    cases = [
        ("#_AQ", True),
        ("#_AQ .title", True),
        ('[id="_AQ"]', True),
        ("[id='_AQ']", True),
        ('div [id="_AQ"] span', True),
        ("  #_AQ  ", True),
        # ↓ 合法的，必须原样放过
        (".search-container", False),
        ("wx-input.search-input", False),
        (".hot-topics .inline-flex", False),
        (".fs-28.text-blacks28.ml-2", False),
        ('[class*="custom-tab-bar"] span', False),
        (".column-card", False),
        ("", False),
        (None, False),
    ]
    bad = ["%r 判成 %s（期望 %s）" % (v, m.looks_like_id_selector(v), want) for v, want in cases
           if m.looks_like_id_selector(v) != want]
    if bad:
        failed.append("按 id 定位的判据：\n      " + "\n      ".join(bad))
        return
    passed += 1
    print("OK   按 id 定位的判据：%d 条正反例（#id / [id=...] 挡住，合法选择器全放过）" % len(cases))


id_selector_rule()


# ---------- 「滑动」每次滑多少：方向可以带幅度（owner 2026/10/14 要的"自己能控制"）----------
# 这几条写成"合规用例"（首步回到首页 + 有个检查），免得体检提示混进断言里
def swipe_case(cid, p2):
    return [
        [cid, "滑动写法", "", 1, "回到首页", "", "", ""],
        [cid, "", "", 2, "滑动", "1", p2, ""],
        [cid, "", "", 3, "检查文字显示", "最新提问", "", ""],
    ]


ok("滑动：老写法（只有方向）= 半屏", swipe_case("S_001", "上"))
ok("滑动：带幅度也通过", swipe_case("S_002", "上 30%"))
ok("滑动：不带空格、全角百分号都认", swipe_case("S_003", "下30％"))
bad_case("滑动：幅度超出范围要报错（不静默夹边界）", [["S_004", "滑", "", 1, "滑动", "1", "上 96%", ""]], "5%~95%")
bad_case("滑动：方向写错要报错", [["S_005", "滑", "", 1, "滑动", "1", "左", ""]], "参数2 只能填")


def swipe_amount_rule():
    """纯函数正反例：合法的幅度解析对，非法的返回 None（交给报错文案去说）"""
    global passed
    cases = [
        ("", ("上", 0.5)),
        ("上", ("上", 0.5)),
        ("下", ("下", 0.5)),
        ("上 30%", ("上", 0.3)),
        ("下30%", ("下", 0.3)),
        ("上30％", ("上", 0.3)),      # 全角百分号
        (" 上 5% ", ("上", 0.05)),    # 边界：最小
        ("下 95%", ("下", 0.95)),     # 边界：最大
        # ↓ 非法的，一律 None
        ("上 4%", None),
        ("上 96%", None),
        ("上 0%", None),
        ("左", None),
        ("上 abc", None),
        ("30%", None),
    ]
    bad = ["%r -> %r（期望 %r）" % (v, m.parse_swipe_amount(v), want) for v, want in cases
           if m.parse_swipe_amount(v) != want]
    if bad:
        failed.append("滑动幅度判据：\n      " + "\n      ".join(bad))
        return
    passed += 1
    # ⚠️ 这行**别**在 %-格式化的字符串里写百分号（`上 30%` 里的 % 会被当成格式符 —— 当天连踩两次）
    print("OK   滑动幅度判据：%d 条正反例（带幅度与不带都认，越界和写错一律挡）" % len(cases))


swipe_amount_rule()
def mixed(name, rows, good_id, bad_id):
    """**这次改动的核心**：一条填错时，别的用例必须一点不受影响。

    正反两面一起验：
      正面 = 好的那条照样导出、步骤齐全、身上没有任何 errors
      反面 = 坏的那条确实被标出来了（不能悄悄放过）
    """
    global passed
    try:
        cases = m.validate(m.read_rows(make(rows)))
    except m.TableError as e:
        failed.append("%s：一条填错不该让整张表报错，实际报了 %s" % (name, e.errors))
        return
    by_id = {c["id"]: c for c in cases}
    if good_id not in by_id or bad_id not in by_id:
        failed.append("%s：导出的用例不全（%s）" % (name, list(by_id)))
        return
    good, bad = by_id[good_id], by_id[bad_id]
    problems = []
    if good.get("errors"):
        problems.append("好的那条也被标了错：%s" % good["errors"])
    if len(good["steps"]) != 2:
        problems.append("好的那条步骤数变成 %d（期望 2）" % len(good["steps"]))
    if not bad.get("errors"):
        problems.append("坏的那条一个错都没标（会被静默放过）")
    if problems:
        failed.append("%s：%s" % (name, "；".join(problems)))
        return
    passed += 1
    print("OK   正反同验 %s（好的照跑，坏的只红它自己）" % name)


mixed(
    "两条用例，只有一条填错",
    [
        ["B_008", "好的那条", "", 1, "回到首页", "", "", ""],
        ["B_008", "", "", 2, "检查文字显示", "积分商城", "", ""],
        ["B_009", "坏的那条", "", 1, "回到首页", "", "", ""],
        ["B_009", "", "", 2, "打开首贝", "", "", ""],
    ],
    "B_008",
    "B_009",
)

# ---------- 表格级反例（归不到哪条用例身上，就该整表停下来）----------
bad_table("空表", [], "一行数据都没有")
bad_table(
    "表头被改",
    [["B_010", "名字", "", 1, "回到首页", "", "", ""]],
    "表头",
    header=["用例编号", "用例名字", "前置条件", "步骤序号", "动作", "参数1", "参数2", "备注"],
)
bad_table(
    "第一行就没编号（上面没得承接）",
    [["", "名字", "", 1, "回到首页", "", "", ""]],
    "「用例编号」是空的",
)


# ---------- 合并单元格（webzdh 式一案一块）----------
def carry(name, rows, want_cases, want_steps):
    """编号空着的行必须**并进上一条用例**，不能凭空多出一条用例、也不能丢步骤。

    为什么要有这条：表里同一条用例的编号/名称被并成了一格，读出来只有第一行有值。
    这里正反两个方向都验：空编号要能承接；表头下面第一行就不能空（没得承接时得报错）。
    """
    global passed
    try:
        cases = m.validate(m.read_rows(make(rows)))
    except m.TableError as e:
        failed.append("%s：本该通过，却整表报错 %s" % (name, e.errors))
        return
    got = [[c["id"], len(c["steps"])] for c in cases]
    want = [[i, n] for i, n in zip(want_cases, want_steps)]
    if got != want:
        failed.append("%s：解析成 %s，期望 %s" % (name, got, want))
        return
    passed += 1
    print("OK   合并排版 %s -> %s" % (name, got))


carry(
    "编号只写首行，下面两行空着",
    [
        ["D_001", "一案多步", "", 1, "回到首页", "", "", ""],
        ["", "", "", 2, "检查文字显示", "积分商城", "", ""],
        ["", "", "", 3, "回到首页", "", "", ""],
    ],
    ["D_001"],
    [3],
)
carry(
    "两块并排，不串味",
    [
        ["D_002", "前一条", "", 1, "回到首页", "", "", ""],
        ["", "", "", 2, "检查文字显示", "积分商城", "", ""],
        ["D_003", "后一条", "", 1, "回到首页", "", "", ""],
    ],
    ["D_002", "D_003"],
    [2, 1],
)


# ---------- 2026/10/14 新增的「前置条件 / 备注」两列（也是按用例并成一格）----------
def case_fields(name, rows, want):
    """一行就能验的四列承接：编号/标题/前置条件/备注 空着 = 沿用上面那条。"""
    global passed
    try:
        cases = m.validate(m.read_rows(make(rows)))
    except m.TableError as e:
        failed.append("%s：本该通过，却整表报错 %s" % (name, e.errors))
        return
    got = [[c["id"], c["name"], c["precondition"], c["note"], len(c["steps"])] for c in cases]
    if got != want:
        failed.append("%s：解析成 %s，期望 %s" % (name, got, want))
        return
    passed += 1
    print("OK   按用例的四列 %s -> %s" % (name, got[0][:4]))


case_fields(
    "前置条件 / 备注 写在首行，下面几行空着",
    [
        ["E_001", "带前置条件", "手机在首页且已登录", 1, "回到首页", "", "", "只读；需求 713000000"],
        ["E_001", "", "", 2, "检查文字显示", "积分商城", "", ""],
        ["E_001", "", "", 3, "回到首页", "", "", ""],
    ],
    [["E_001", "带前置条件", "手机在首页且已登录", "只读；需求 713000000", 3]],
)
case_fields(
    "前置条件 / 备注 补写在中间的步骤行上（合并单元格没并到位时不算错）",
    [
        ["E_002", "后补的", "", 1, "回到首页", "", "", ""],
        ["E_002", "", "要登录", 2, "检查文字显示", "积分商城", "", "只读"],
    ],
    [["E_002", "后补的", "要登录", "只读", 2]],
)
case_fields(
    "两条用例各带各的前置条件，不串味",
    [
        ["E_003", "前一条", "前置A", 1, "回到首页", "", "", "备注A"],
        ["E_003", "", "", 2, "检查文字显示", "积分商城", "", ""],
        ["E_004", "后一条", "前置B", 1, "回到首页", "", "", "备注B"],
    ],
    [["E_003", "前一条", "前置A", "备注A", 2], ["E_004", "后一条", "前置B", "备注B", 1]],
)


def case_field_conflict(name, rows, want):
    """同一条用例的「备注」（或前置条件）写了两个不一样的值 -> 报错，不能静默丢一个。"""
    global passed
    try:
        cases = m.validate(m.read_rows(make(rows)))
    except m.TableError as e:
        failed.append("%s：本该把错误归到用例身上，却整表报错 %s" % (name, e.errors))
        return
    msgs = [msg for c in cases for msg in (c.get("errors") or [])]
    if any(want in msg for msg in msgs):
        passed += 1
        print("OK   按用例的四列 %s -> %s" % (name, msgs[0][:48]))
        return
    failed.append("%s：没报出「%s」，实际是 %s" % (name, want, msgs or "（一条错误都没有）"))


case_field_conflict(
    "同一条用例的备注写了两个值",
    [
        ["E_005", "备注冲突", "", 1, "回到首页", "", "", "第一个备注"],
        ["E_005", "", "", 2, "检查文字显示", "积分商城", "", "第二个备注"],
    ],
    "「备注」已经是",
)
case_field_conflict(
    "同一条用例的前置条件写了两个值",
    [
        ["E_006", "前置条件冲突", "要登录", 1, "回到首页", "", "", ""],
        ["E_006", "", "不用登录", 2, "检查文字显示", "积分商城", "", ""],
    ],
    "「前置条件」已经是",
)
case_field_conflict(
    "同一条用例的标题写了两个值（老规矩，仍然要抓）",
    [
        ["E_007", "标题甲", "", 1, "回到首页", "", "", ""],
        ["E_007", "标题乙", "", 2, "检查文字显示", "积分商城", "", ""],
    ],
    "「用例标题」已经是",
)


def actions_sheet_at_top(times):
    """「动作清单」刷 N 次也必须**一直待在第 1 行**，而且下拉框指的那几行不是空的。

    这是 2026/10/13 真踩到的 bug 的反向验证：以前 append 是往 max_row 后面接，
    每刷一次清单就往下挪一截（刷 9 次挪到第 436 行），而下拉框还指着 A2:A48 —— 空行，
    于是下拉框里一个动作都选不出来。
    """
    global passed
    wb = Workbook()
    ws = wb.active
    ws.title = m.SHEET_ACTIONS
    for _ in range(times):
        m.fill_actions_sheet(ws)
    name = "动作清单刷 %d 次仍在表头下方" % times
    if ws.cell(1, 1).value != "动作":
        failed.append("%s：第 1 行是 %r，不是表头（清单被 append 挤下去了）" % (name, ws.cell(1, 1).value))
        return
    got = [ws.cell(r, 1).value for r in range(2, len(m.ACTIONS) + 2)]
    # 顺序必须等于**按用途分组**之后的顺序（表格下拉框读起来像目录，就是靠这一条）
    want = [a["name"] for a in m.ordered_actions()]
    if got != want:
        failed.append("%s：清单内容对不上（%s…）" % (name, got[:3]))
        return
    if ws.cell(len(m.ACTIONS) + 2, 1).value is not None:
        failed.append("%s：清单下面还留着旧内容（第 %d 行）" % (name, len(m.ACTIONS) + 2))
        return
    passed += 1
    print("OK   动作清单刷 %d 次：%d 个动作都在第 2~%d 行" % (times, len(want), len(want) + 1))


actions_sheet_at_top(1)
actions_sheet_at_top(3)
actions_sheet_at_top(9)  # 就是踩坑那次刷的次数


def example_params(action):
    """把「怎么填（示例）」拆成参数1 / 参数2（有两个参数的动作写成 `参数1 / 参数2`）"""
    ex = str(m.ACTION_EXAMPLES.get(action["name"], "")).strip()
    if not action["p1"]:
        return "", ""
    # 用「 / 」（带空格）当分隔符：示例本身可能带括号说明，括号里也可能有斜杠
    parts = [p.strip() for p in ex.split(" / ")] if action["p2"] and " / " in ex else [ex]
    parts = [re.sub(r"（[^）]*）\s*$", "", p).strip() for p in parts]
    return parts[0], (parts[1] if len(parts) > 1 else "")


def examples_are_copyable():
    """**每条示例都必须"照抄就能过校验"** —— 把 49 条示例当成参数填一张表，跑一遍导出校验。

    为什么值得单独验：示例是给人照抄的。示例自己写错（比如填了个不存在的候选值、或该填的参数空着），
    照抄的人会被导出报错挡住 —— 那时他只会觉得"文档骗人"。所以让闸门替他们先抄一遍。
    """
    global passed
    rows = []
    for a in m.ACTIONS:
        p1, p2 = example_params(a)
        rows.append(
            ["EX_%s" % a["name"], "示例校验", "手机在测试环境", 1, a["name"], p1, p2, ""]
        )
    try:
        cases = m.validate(m.read_rows(make(rows)))
    except m.TableError as e:
        failed.append("示例校验：本该通过，整表却报错 %s" % e.errors)
        return
    bad = [(c["id"], c["errors"]) for c in cases if c.get("errors")]
    if bad:
        failed.append("示例校验：这些示例照抄会被报错 %s" % bad[:3])
        return
    if len(cases) != len(m.ACTIONS):
        failed.append("示例校验：用例数 %d 与动作数 %d 对不上" % (len(cases), len(m.ACTIONS)))
        return
    passed += 1
    print("OK   示例校验：%d 个动作的示例照抄都能过导出校验" % len(cases))


examples_are_copyable()


def refresh_is_idempotent():
    """刷新两次：**第二次不许改文件**。

    为什么盯它：`--dropdown` 每次都 `wb.save()` 重写整个 xlsx，哪怕内容一点没变 ——
    文件字节也会变，于是 git 里不停冒出"没有实际改动"的二进制 diff（真踩到：想确认下拉框刷没刷，
    结果工作区又脏了一次）。现在刷完先比对，没变就不写；这条自检把那个行为钉住。
    """
    global passed
    import hashlib
    from pathlib import Path

    xlsx = os.path.join(TMP, "refresh_idempotent.xlsx")
    jpath = os.path.join(TMP, "refresh_idempotent.json")
    m.build_template(Path(xlsx))
    digest = lambda p: hashlib.md5(open(p, "rb").read()).hexdigest()

    # 让"候选名字"取不到（指向不存在的配置）——自检不该依赖测试库；取不到会优雅降级
    backup = os.environ.get("UITEST_DB_CONFIG")
    os.environ["UITEST_DB_CONFIG"] = os.path.join(TMP, "no_such_db.yaml")
    try:
        m.main(["--dropdown", "--xlsx", xlsx, "--json", jpath])  # 第一次：模板 → 补候选值，会写
        h1 = digest(xlsx)
        m.main(["--dropdown", "--xlsx", xlsx, "--json", jpath])  # 第二次：什么都没变，不该写
        h2 = digest(xlsx)
    finally:
        if backup is None:
            os.environ.pop("UITEST_DB_CONFIG", None)
        else:
            os.environ["UITEST_DB_CONFIG"] = backup

    if h1 != h2:
        failed.append("刷新幂等性：第二次刷新也重写了文件（会造成无意义的 git diff）")
        return
    passed += 1
    print("OK   刷新幂等：内容没变就不重写文件（刷两次，第二次文件字节不变）")


refresh_is_idempotent()


def optional_flags_reach_the_json():
    """「可不填」的参数，必须把 p1_optional / p2_optional **带进 json**。

    为什么单独验这条（2026/10/14 真踩到）："这个参数是不是必填"在两侧各判断一次 ——
    Python 侧（导出校验）和 JS 侧（真机执行前再查一遍）。加「点文字」的「第几个（可不填）」时，
    只有 Python 侧认这个标记，JS 侧还在按"这一格有值吗"判断：于是表格里参数2 留空，
    真机跑起来被当成漏填，**用例直接报错**。当时的自检只覆盖了 Python 侧，没拦住。
    现在这条自检盯的是"标记有没有传到执行端"，正反两个方向都验：
      · 写着「可不填」的参数 -> 标记必须是 True（否则 JS 会误判成漏填）
      · 没写「可不填」的参数 -> 标记必须是 False（否则真漏填也会被放过）
    """
    global passed
    import json as _json
    from pathlib import Path

    xlsx = os.path.join(TMP, "optional.xlsx")
    jpath = os.path.join(TMP, "optional.json")
    m.build_template(Path(xlsx))
    m.main(["--xlsx", xlsx, "--json", jpath])
    data = _json.loads(open(jpath, encoding="utf-8").read())

    problems = []
    for a in data["actions"]:
        for f in ("p1", "p2"):
            label = str(a.get(f) or "")
            wanted = "可不填" in label
            got = bool(a.get(f + "_optional"))
            if wanted != got:
                problems.append(
                    "动作「%s」的 %s：文案=%r，标记 %s_optional=%s（期望 %s）"
                    % (a["name"], f, label, f, got, wanted)
                )
    if problems:
        failed.append("可选参数标记：\n      " + "\n      ".join(problems))
        return

    # 反向：必须真的存在至少一个可选参数（否则这条自检在空跑）
    n_opt = sum(1 for a in data["actions"] if a.get("p1_optional") or a.get("p2_optional"))
    if n_opt == 0:
        failed.append("可选参数标记：一个可选参数都没有 —— 这条自检在空跑，检查 ACTIONS 里的 p?_optional")
        return
    passed += 1
    print("OK   可选参数标记：%d 个动作带「可不填」参数，标记与文案一致（会传到执行端）" % n_opt)


optional_flags_reach_the_json()


# ---------- 「可不填」的参数（2026/10/14 随「第几个」一起加的概念）----------
# 为什么单独验：所有参数以前都是"声明了就必须填"。加了可选参数之后，
# **老用例（参数2 留空）必须照常通过**，而不是被新规则误伤 —— 这条正是自检当场抓到的。
ok(
    "可选参数留空 -> 走默认值（老用例不受影响）",
    [
        ["F_001", "留空", "", 1, "回到首页", "", "", ""],
        ["F_001", "", "", 2, "点选择器", ".hot-topics .inline-flex", "", ""],
        ["F_001", "", "", 3, "检查文字显示", "最新提问", "", ""],
    ],
)
ok(
    "可选参数填了 -> 也通过",
    [
        ["F_002", "填了", "", 1, "回到首页", "", "", ""],
        ["F_002", "", "", 2, "点选择器", ".hot-topics .inline-flex", "2", ""],
        ["F_002", "", "", 3, "检查文字显示", "最新提问", "", ""],
    ],
)
bad_case(
    "可选参数填了不合法的值 -> 仍然报错",
    [
        ["F_003", "值不对", "", 1, "回到首页", "", "", ""],
        ["F_003", "", "", 2, "点选择器", ".hot-topics .inline-flex", "第二个", ""],
    ],
    "「第几个」只能填",
)
bad_case(
    "可选参数填 0 -> 报错（从 1 数起）",
    [
        ["F_004", "零", "", 1, "回到首页", "", "", ""],
        ["F_004", "", "", 2, "点选择器", ".hot-topics .inline-flex", "0", ""],
    ],
    "「第几个」只能填",
)


def merge_keeps_content():
    """并格**不能改内容**：并格前后导出来的用例（编号/标题/前置条件/备注/步骤数）必须一模一样。

    这是 2026/10/13 那条真 bug 的反向验证：一开始把「编号/名称」并成了一整块 A:B，
    而 openpyxl 合并时会把除左上角以外的单元格清空 —— 名字正好在 B 列，被吃掉，
    导出时用例名静默退化成编号（看着"跑通了"，其实用例名没了）。
    2026/10/14 加了「前置条件」并把「备注」也改成按用例合并，同一个坑又多踩的机会：
    四列都得**一列一列地并**，并完内容一个字不能少。
    """
    global passed
    from openpyxl import load_workbook

    wb = Workbook()
    ws = wb.active
    ws.title = m.SHEET_DATA
    ws.append(m.HEADERS)
    for row in m.SAMPLE_ROWS:
        ws.append(row)
    wb.save(PATH)
    snapshot = lambda: [
        [c["id"], c["name"], c["precondition"], c["note"], len(c["steps"])]
        for c in m.validate(m.read_rows(PATH))
    ]
    before = snapshot()
    merged = 0

    wb = load_workbook(PATH)
    merged = m.merge_case_cells(wb[m.SHEET_DATA])
    wb.save(PATH)
    after = snapshot()

    if before != after:
        failed.append("并格改了内容：并之前 %s，并之后 %s" % (before, after))
        return
    if any(name == cid for cid, name, _p, _no, _n in after):
        failed.append("并格之后有用例名退化成了编号：%s" % after)
        return
    # 并完再并一次：块数必须一样（重复执行不叠加、不越并）
    wb = load_workbook(PATH)
    again = m.merge_case_cells(wb[m.SHEET_DATA])
    if again != merged:
        failed.append("并格不幂等：第一次并了 %d 块，第二次又并了 %d 块" % (merged, again))
        return
    passed += 1
    print("OK   并格不改内容（%d 条用例 / %d 块，四列都并了，名字与备注都在）" % (len(after), merged))


merge_keeps_content()

# ---------- 体检（lint）反例：只提示，不拦截 ----------
ok(
    "一条用例没有任何检查",
    [
        ["C_001", "无检查", "", 1, "回到首页", "", "", ""],
        ["C_001", "", "", 2, "切到用户中心", "", "", ""],
        ["C_001", "", "", 3, "回到首页", "", "", ""],
    ],
    want_warn_count=1,
)
ok(
    "第一步不回到首页",
    [
        ["C_002", "起点不对", "", 1, "检查文字显示", "积分商城", "", ""],
        ["C_002", "", "", 2, "回到首页", "", "", ""],
    ],
    want_warn_count=1,
)
ok(
    "最后一步停在二级页",
    [
        ["C_003", "结尾不对", "", 1, "回到首页", "", "", ""],
        ["C_003", "", "", 2, "检查文字显示", "热门专栏", "", ""],
        ["C_003", "", "", 3, "打开搜索并输入", "test", "", ""],
    ],
    want_warn_count=1,
)
ok("三条规矩全犯", [["C_004", "全犯", "", 1, "打开答题活动", "", "", ""]], want_warn_count=3)

print()
if failed:
    print("有 %d 项不符合预期：" % len(failed))
    for f in failed:
        print("  · %s" % f)
    sys.exit(1)
print("全部通过（%d 项）" % passed)
